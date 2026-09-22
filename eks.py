import os
import json
import pandas as pd
import docx
from google import genai
from openpyxl.styles import Font, PatternFill, Alignment

# gemini
client = genai.Client(api_key="api key")

try:
    import win32com.client as win32
except ImportError:
    win32 = None
    print("Peringatan: Library 'pywin32' belum terinstall. Fitur konversi otomatis .doc ke .docx dimatikan.")
    print("Silakan install melalui CMD dengan perintah: pip install pywin32\n")

def konversi_doc_ke_docx(folder_path):
    """Fungsi untuk mencari file .doc dan mengonversinya menjadi .docx"""
    if win32 is None:
        return 

    doc_files = [f for f in os.listdir(folder_path) if f.lower().endswith(".doc") and not f.startswith("~")]
    
    if not doc_files:
        return 
        
    print(f"🔄 Mendeteksi {len(doc_files)} file berformat .doc lama. Memulai konversi ke .docx...")
    
    try:
        word = win32.Dispatch('Word.Application')
        word.Visible = False
        
        for filename in doc_files:
            in_file = os.path.abspath(os.path.join(folder_path, filename))
            out_filename = filename + "x" 
            out_file = os.path.abspath(os.path.join(folder_path, out_filename))

            if not os.path.exists(out_file):
                print(f"   -> Mengonversi: {filename}")
                try:
                    doc = word.Documents.Open(in_file)
                    doc.SaveAs(out_file, FileFormat=16)
                    doc.Close()
                except Exception as e:
                    print(f"Gagal mengonversi {filename}: {e}")
                    
        word.Quit()
        print("Konversi selesai.\n")
        
    except Exception as e:
        print(f"Gagal memanggil sistem Microsoft Word: {e}\n")

def proses_teks_dengan_ai(teks_mentah):
    """Fungsi untuk mengirim teks RPS ke Gemini dan memintanya mengembalikan JSON"""
    
    prompt = """
    Anda adalah asisten data scraper. Berikut adalah isi dokumen Rencana Pembelajaran Semester (RPS).
    Tugas Anda:
    1. Cari dan ekstrak pokok bahasan / materi pembelajaran untuk setiap pertemuan/minggu.
    2. PENTING: JANGAN meringkas materi. Ekstrak seluruh sub-bab, poin-poin detail, dan penjelasan materi secara UTUH dan LENGKAP persis seperti yang tertulis di dokumen aslinya.
    3. Kembalikan data HANYA dalam format JSON array yang valid, tanpa teks awalan/akhiran, tanpa markdown blok (```json).
    Struktur JSON yang wajib digunakan:
    [
      {
        "pertemuan": "1",
        "materi_id": "Teks materi bahasa Indonesia"
      }
    ]
    
    Teks Dokumen:
    """ + teks_mentah

    try:
        response = client.models.generate_content(
            model='gemini-3.5-flash',
            contents=prompt
        )
        hasil_teks = response.text.strip()
        
        if hasil_teks.startswith("```json"):
            hasil_teks = hasil_teks.replace("```json", "", 1)
        if hasil_teks.startswith("```"):
            hasil_teks = hasil_teks.replace("```", "", 1)
        if hasil_teks.endswith("```"):
            hasil_teks = hasil_teks[::-1].replace("```", "", 1)[::-1]
            
        return json.loads(hasil_teks.strip())
    except Exception as e:
        print(f"\n   -> [Error AI]: {e}")
        return []


def ekstrak_materi_rps(folder_path):
    konversi_doc_ke_docx(folder_path)
    
    print("MEMULAI EKSTRAKSI DOKUMEN DENGAN AI...")
    
    all_data = []
    
    files = [f for f in os.listdir(folder_path) if f.endswith(".docx") and not f.startswith("~")]
    total_files = len(files)

    for idx, filename in enumerate(files, 1):
        file_path = os.path.join(folder_path, filename)
        print(f"[{idx}/{total_files}] Membaca dengan AI: {filename[:25]}... ", end="")
        
        try:
            doc = docx.Document(file_path)
            nama_mk = filename.replace(".docx", "").strip()
            
            teks_gabungan = []
            for p in doc.paragraphs:
                if p.text.strip(): teks_gabungan.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    baris_teks = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if baris_teks: teks_gabungan.append(" | ".join(baris_teks))
            
            teks_dokumen_full = "\n".join(teks_gabungan)
            
            if len(teks_dokumen_full) > 40000:
                teks_dokumen_full = teks_dokumen_full[:40000]

            data_json = proses_teks_dengan_ai(teks_dokumen_full)
            
            if not data_json:
                print("GAGAL (AI tidak menemukan materi atau terjadi error)")
                continue

            for item in data_json:
                all_data.append({
                    "Kode Mata Kuliah": nama_mk,
                    "Pertemuan": item.get("pertemuan", ""),
                    "Materi Indonesia": item.get("materi_id", ""),
                    "Materi Inggris": item.get("materi_en", "")
                })
            
            for _ in range(3):
                all_data.append({"Kode Mata Kuliah": "", "Pertemuan": "", "Materi Indonesia": "", "Materi Inggris": ""})
                
            print(f" Selesai ({len(data_json)} Pertemuan Diekstrak)")
            
        except Exception as e:
            print(f"ERROR: {e}")

    if all_data:
        print("\nMenyusun dan merapikan file Excel...")
        df = pd.DataFrame(all_data)
        output_file = os.path.join(folder_path, "Extrak_output_AI.xlsx")
        
        try:
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Data_RPS')
                worksheet = writer.sheets['Data_RPS']
                
                header_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
                header_font = Font(color="FFFFFF", bold=True)
                header_align = Alignment(horizontal="center", vertical="center")
                
                for col in range(1, 5):
                    cell = worksheet.cell(row=1, column=col)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_align
                
                worksheet.column_dimensions['A'].width = 30 
                worksheet.column_dimensions['B'].width = 15 
                worksheet.column_dimensions['C'].width = 75 
                worksheet.column_dimensions['D'].width = 75 
                
                wrap_alignment = Alignment(wrap_text=True, vertical="top")
                for row in worksheet.iter_rows(min_row=2, max_row=len(all_data)+1, min_col=1, max_col=4):
                    for cell in row:
                        cell.alignment = wrap_alignment

            print(f"File Excel selesai dibuat, buka di:\n {output_file}")
            
        except PermissionError:
            print("\n GAGAL MENYIMPAN: File Excel sedang terbuka! Tutup dulu file Excel-nya lalu jalankan ulang.")
    else:
        print("\n Tidak ada data yang berhasil diekstrak.")

LOKASI_FOLDER = r"C:\Users\Sofiandi\Downloads\RPS_FOLDER" 
ekstrak_materi_rps(LOKASI_FOLDER)
